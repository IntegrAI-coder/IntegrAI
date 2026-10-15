#!/usr/bin/env python3
"""Convierte la rúbrica oficial de Integra (.docx) en el marco JSON que usa el sistema.

    python tools/importar_rubrica.py "Rúbrica progresión Identidad y Autonomía.docx"
    → privado/marcos/integra_identidad_autonomia.json   (NO se versiona ni se publica: material de Integra)

El docx trae 3 tablas (Sala Cuna, Nivel Medio, Transición) de 8 filas:
  "N. Nombre" | descripción del nivel | "• Conducta: … • Ejemplo: … • Ejemplo: …"
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "privado" / "marcos" / "integra_identidad_autonomia.json"   # fuera de data/ (el frontend la publica)

# Tramos oficiales (ChileAtiende, ficha 1786): edad al 31 de marzo del año escolar
TRAMOS = [
    {"id": "sala_cuna", "nombre": "Sala Cuna", "sigla": "SC", "edad_min_meses": 0, "edad_max_meses": 24,
     "subtramos": [{"id": "sala_cuna_menor", "nombre": "Sala Cuna Menor", "edad_min_meses": 0, "edad_max_meses": 12},
                   {"id": "sala_cuna_mayor", "nombre": "Sala Cuna Mayor", "edad_min_meses": 12, "edad_max_meses": 24}]},
    {"id": "nivel_medio", "nombre": "Nivel Medio", "sigla": "NM", "edad_min_meses": 24, "edad_max_meses": 48,
     "subtramos": [{"id": "medio_menor", "nombre": "Medio Menor", "edad_min_meses": 24, "edad_max_meses": 36},
                   {"id": "medio_mayor", "nombre": "Medio Mayor", "edad_min_meses": 36, "edad_max_meses": 48}]},
    {"id": "transicion", "nombre": "Transición", "sigla": "TR", "edad_min_meses": 48, "edad_max_meses": 72,
     "subtramos": [{"id": "nt1", "nombre": "Primer Nivel de Transición", "edad_min_meses": 48, "edad_max_meses": 60},
                   {"id": "nt2", "nombre": "Segundo Nivel de Transición", "edad_min_meses": 60, "edad_max_meses": 72}]},
]
NOMBRES = ["Comenzando", "Explorando", "Desarrollando", "Afianzando", "Integrando", "Ampliando", "Consolidando", "Profundizando"]


def _items(texto: str) -> tuple[list[str], list[str]]:
    """'• Conducta: a • Ejemplo: b • Ejemplo: c' -> (['a'], ['b', 'c'])."""
    conductas, ejemplos = [], []
    for parte in re.split(r"[•\n]+", texto):
        parte = parte.strip()
        m = re.match(r"(Conducta|Ejemplo)\s*:\s*(.+)", parte, re.S)
        if m:
            (conductas if m.group(1) == "Conducta" else ejemplos).append(m.group(2).strip())
    return conductas, ejemplos


def importar(docx_path: str) -> dict:
    import docx  # python-docx
    d = docx.Document(docx_path)
    if len(d.tables) != 3:
        sys.exit(f"se esperaban 3 tablas (una por tramo), hay {len(d.tables)}")
    tramos = []
    for tramo, tabla in zip(TRAMOS, d.tables):
        niveles = []
        for fila in tabla.rows[1:]:
            celdas = [c.text.strip() for c in fila.cells]
            m = re.match(r"(\d)\.\s*(\w+)", celdas[0])
            n, nombre = int(m.group(1)), m.group(2)
            assert NOMBRES[n - 1] == nombre, f"nivel {n}: '{nombre}' no coincide con '{NOMBRES[n - 1]}'"
            conductas, ejemplos = _items(celdas[2])
            niveles.append({"nivel": n, "nombre": nombre, "descripcion": celdas[1],
                            "conductas": conductas, "ejemplos": ejemplos})
        assert [x["nivel"] for x in niveles] == list(range(1, 9)), f"{tramo['id']}: niveles incompletos"
        tramos.append({**tramo, "niveles": niveles})
    return {
        "marco": "integra_iya", "version": "1.0",
        "nucleo": {"id": "identidad_autonomia", "nombre": "Identidad y Autonomía"},
        "ambito": "Desarrollo Personal y Social",
        "fuente": {"archivo": Path(docx_path).name,
                   "sha256": hashlib.sha256(Path(docx_path).read_bytes()).hexdigest(),
                   "origen": "Fundación Integra (material entregado para la hackathon; no publicar)"},
        "regla_tramo": "edad al 31 de marzo del año escolar (ChileAtiende, ficha 1786)",
        "tramos": tramos,
    }


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    marco = importar(sys.argv[1])
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    SALIDA.write_text(json.dumps(marco, ensure_ascii=False, indent=1), encoding="utf-8")
    n = sum(len(t["niveles"]) for t in marco["tramos"])
    a = sum(1 + len(l["conductas"]) + len(l["ejemplos"]) for t in marco["tramos"] for l in t["niveles"])
    print(f"ok: {len(marco['tramos'])} tramos, {n} niveles, {a} anclas -> {SALIDA}")
