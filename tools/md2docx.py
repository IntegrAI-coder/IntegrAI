#!/usr/bin/env python3
"""Convierte nuestro Markdown a .docx (Google Doc-friendly): títulos, tablas, listas, código."""
from __future__ import annotations
import re, sys
from docx import Document
from docx.shared import Pt

def add_runs(par, text: str) -> None:
    for p in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not p:
            continue
        if p.startswith("**") and p.endswith("**"):
            par.add_run(p[2:-2]).bold = True
        elif p.startswith("`") and p.endswith("`"):
            r = par.add_run(p[1:-1]); r.font.name = "Consolas"
        else:
            par.add_run(p)

def convert(md_path: str, out_path: str) -> None:
    lines = open(md_path, encoding="utf-8").read().splitlines()
    doc = Document()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            i += 1; buf = []
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i]); i += 1
            p = doc.add_paragraph(); r = p.add_run("\n".join(buf)); r.font.name = "Consolas"; r.font.size = Pt(9)
            i += 1; continue
        if re.match(r"^\|.*\|$", line) and i + 1 < len(lines) and re.match(r"^\|[-: |]+\|$", lines[i + 1]):
            header = [c.strip() for c in line.strip("|").split("|")]
            i += 2; rows = []
            while i < len(lines) and re.match(r"^\|.*\|$", lines[i]):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")]); i += 1
            t = doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                c = t.rows[0].cells[j]; c.text = ""; add_runs(c.paragraphs[0], h)
            for r_ in rows:
                cells = t.add_row().cells
                for j, cc in enumerate(r_[:len(header)]):
                    cells[j].text = ""; add_runs(cells[j].paragraphs[0], cc)
            doc.add_paragraph(); continue
        if line.startswith("### "): doc.add_heading(line[4:], level=3)
        elif line.startswith("## "): doc.add_heading(line[3:], level=2)
        elif line.startswith("# "): doc.add_heading(line[2:], level=1)
        elif line.startswith("> "):
            p = doc.add_paragraph(); add_runs(p, line[2:])
            for r in p.runs:
                r.italic = True
        elif re.match(r"^[-*] ", line):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number"); add_runs(p, re.sub(r"^\d+\. ", "", line))
        elif line.strip() in ("", "---"):
            pass
        else:
            p = doc.add_paragraph(); add_runs(p, line)
        i += 1
    doc.save(out_path)

if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
    print("docx:", sys.argv[2])
